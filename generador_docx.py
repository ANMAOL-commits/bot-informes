"""Genera el informe en Word (.docx) con el formato del modelo, usando python-docx."""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

FUENTE = "Arial"
TOKEN = re.compile(r"(\*\*.+?\*\*|\[\[.+?\]\])")
ALTO_MAX_FOTO_CM = 12


# --------------------------------------------------------------------------- utilidades
def _runs(par, texto, size=None, bold=False, italic=False):
    """Escribe texto con **negrita** y [[texto a verificar]] (resaltado amarillo)."""
    for seg in filter(None, TOKEN.split(texto or "")):
        if seg.startswith("**") and seg.endswith("**") and len(seg) > 4:
            _runs(par, seg[2:-2], size, True, italic)
            continue
        resaltar = seg.startswith("[[") and seg.endswith("]]") and len(seg) > 4
        r = par.add_run(seg[2:-2] if resaltar else seg)
        r.bold = True if bold else None
        r.italic = True if italic else None
        if size:
            r.font.size = Pt(size)
        if resaltar:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _sin_tema(estilo):
    """Quita los atributos de fuente 'de tema' para que Word respete Arial."""
    rpr = estilo.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FUENTE)


def _sombrear(celda, color):
    tcpr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcpr.append(shd)


def _borde_superior(par):
    ppr = par._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    for k, v in (("val", "single"), ("sz", "4"), ("space", "8"), ("color", "BFBFBF")):
        top.set(qn(f"w:{k}"), v)
    pbdr.append(top)
    ppr.insert_element_before(
        pbdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
        "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
        "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
        "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
    )


def _campo(par, instruccion, size=8):
    r = par.add_run()
    r.bold = True
    r.font.size = Pt(size)
    for tipo, texto in (("begin", None), ("instr", f" {instruccion} "), ("separate", None), ("txt", "1"), ("end", None)):
        if tipo == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = texto
        elif tipo == "txt":
            el = OxmlElement("w:t")
            el.text = texto
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        r._r.append(el)


def _texto_celda(celda, texto, size=8, bold=False, centrado=False):
    p = celda.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _runs(p, texto, size=size, bold=bold)
    celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return p


def _quitar_primer_parrafo(contenedor):
    p = contenedor.paragraphs[0]._p
    p.getparent().remove(p)


# --------------------------------------------------------------------------- encabezado y pie
def _encabezado(doc, spec, logo):
    sec = doc.sections[0]
    sec.header.is_linked_to_previous = False
    h = sec.header
    anchos = [Cm(3.88), Cm(7.24), Cm(3.88)]
    t = h.add_table(4, 3, Cm(15))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for fila in t.rows:
        for i, c in enumerate(fila.cells):
            c.width = anchos[i]

    celda_logo = t.cell(0, 0).merge(t.cell(3, 0))
    celda_logo.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = celda_logo.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    if logo and Path(logo).exists():
        p.add_run().add_picture(str(logo), width=Cm(2.6))
    else:
        r = p.add_run("ENFOQUE")
        r.bold = True
        r.font.size = Pt(11)
        p2 = celda_logo.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        p2.add_run("INGENIERÍA").font.size = Pt(9)

    arriba = t.cell(0, 1).merge(t.cell(1, 1))
    abajo = t.cell(2, 1).merge(t.cell(3, 1))
    _texto_celda(arriba, spec.get("tipo_label", "INFORME TÉCNICO"), size=10, bold=True, centrado=True)
    _texto_celda(abajo, spec.get("subtitulo", ""), size=10, bold=True, centrado=True)

    _texto_celda(t.cell(0, 2), f"CÓDIGO {spec.get('codigo', 'DOC')}")
    _texto_celda(t.cell(1, 2), f"REVISIÓN {spec.get('revision', '00')}")
    _texto_celda(t.cell(2, 2), spec.get("fecha", ""))
    p = _texto_celda(t.cell(3, 2), "Página ")
    _campo(p, "PAGE")
    p.add_run(" de ").font.size = Pt(8)
    _campo(p, "NUMPAGES")

    h.add_paragraph().paragraph_format.space_after = Pt(0)
    _quitar_primer_parrafo(h)


def _pie(doc):
    sec = doc.sections[0]
    sec.footer.is_linked_to_previous = False
    f = sec.footer
    t = f.add_table(1, 3, Cm(15))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.rows[0].height = Cm(0.6)
    for i, etiqueta in enumerate(("Elaboró:", "Revisó:", "Aprobó:")):
        t.cell(0, i).width = Cm(5)
        _texto_celda(t.cell(0, i), etiqueta, size=8, bold=True)
    f.add_paragraph().paragraph_format.space_after = Pt(0)
    _quitar_primer_parrafo(f)


# --------------------------------------------------------------------------- cuerpo
def _imagen(doc, carpeta, ruta, epigrafe, ancho_cm):
    archivo = Path(ruta) if Path(ruta).is_absolute() else Path(carpeta) / ruta
    if not archivo.exists():
        p = doc.add_paragraph()
        _runs(p, f"[[Imagen no encontrada: {ruta}]]")
        return
    with Image.open(archivo) as im:
        w, h = im.size
    ancho = min(ancho_cm, ALTO_MAX_FOTO_CM * w / h)
    doc.add_picture(str(archivo), width=Cm(ancho))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    if epigrafe:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.paragraph_format.space_after = Pt(10)
        _runs(c, epigrafe, size=9, italic=True)


def _tabla(doc, bloque):
    enc = bloque.get("header", [])
    filas = bloque.get("rows", [])
    n = max(len(enc), 1)
    anchos = bloque.get("widths") or [8504 // n] * n
    anchos_cm = [a / 567 for a in anchos]
    t = doc.add_table(rows=1 + len(filas), cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, titulo in enumerate(enc):
        c = t.cell(0, i)
        c.width = Cm(anchos_cm[i])
        _texto_celda(c, str(titulo), size=9, bold=True)
        _sombrear(c, "E7E6E6")
    trpr = t.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader"))
    for r, fila in enumerate(filas, 1):
        for i in range(n):
            c = t.cell(r, i)
            c.width = Cm(anchos_cm[i])
            _texto_celda(c, str(fila[i]) if i < len(fila) and fila[i] is not None else "", size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def construir_docx(spec, carpeta, salida, logo=None):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(3.0)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.header_distance = sec.footer_distance = Cm(1.25)

    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    _sin_tema(normal)
    h1 = doc.styles["Heading 1"]
    h1.font.name = FUENTE
    h1.font.size = Pt(11)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0, 0, 0)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    _sin_tema(h1)
    doc.styles["List Bullet"].paragraph_format.space_after = Pt(3)

    _encabezado(doc, spec, logo)
    _pie(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run((spec.get("titulo") or "").upper())
    r.bold = True
    r.font.size = Pt(11)

    for n, s in enumerate(spec.get("secciones", []), 1):
        h = doc.add_paragraph(style="Heading 1")
        h.paragraph_format.keep_with_next = True
        h.add_run(f"{n}. {s.get('titulo', '')}")
        _borde_superior(h)
        for b in s.get("bloques", []):
            tipo = b.get("type")
            if tipo == "p":
                _runs(doc.add_paragraph(), b.get("text", ""))
            elif tipo == "bullets":
                for it in b.get("items", []):
                    _runs(doc.add_paragraph(style="List Bullet"), it)
            elif tipo == "numbered":
                for i, it in enumerate(b.get("items", []), 1):
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(1.27)
                    p.paragraph_format.first_line_indent = Cm(-0.63)
                    p.paragraph_format.space_after = Pt(4)
                    p.paragraph_format.tab_stops.add_tab_stop(Cm(1.27))
                    p.add_run(f"{i}.\t")
                    _runs(p, it)
            elif tipo == "table":
                _tabla(doc, b)
            elif tipo == "figure":
                _imagen(doc, carpeta, b.get("path", ""), b.get("caption"), b.get("widthCm", 14))

    anexo = spec.get("anexo")
    if anexo and anexo.get("fotos"):
        doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        p.add_run(f"ANEXO: {anexo.get('titulo', 'Registro fotográfico')}")
        for i, f in enumerate(anexo["fotos"], 1):
            _imagen(doc, carpeta, f.get("path", ""), f.get("caption") or f"Foto {i}", f.get("widthCm", 11))

    doc.core_properties.author = "Enfoque Ingeniería"
    doc.core_properties.title = spec.get("titulo", "Informe")
    doc.save(str(salida))
